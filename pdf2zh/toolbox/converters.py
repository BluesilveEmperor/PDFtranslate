import io
import os

import pymupdf as fitz

from ._util import ensure_parent, parse_pages


def _open(path):
    if not os.path.isfile(path):
        raise FileNotFoundError(f"文件不存在: {path}")
    return fitz.open(path)


def pdf_to_image(input, outdir, dpi=150, fmt="png", pages=None, quality=90):
    from PIL import Image

    os.makedirs(outdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(input))[0]
    outputs = []
    with _open(input) as doc:
        sel = parse_pages(pages, doc.page_count)
        for p in sel:
            page = doc[p]
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            mode = "RGB"
            data = pix.samples
            img = Image.frombytes(mode, (pix.width, pix.height), data)
            name = f"{stem}_{p + 1:04d}.{fmt}"
            path = os.path.join(outdir, name)
            if fmt in ("jpg", "jpeg"):
                img.save(path, "JPEG", quality=quality)
            else:
                img.save(path, fmt.upper())
            outputs.append(path)
    return outputs


def image_to_pdf(images, output, page_size="auto", margin=0):
    from PIL import Image

    ensure_parent(output)
    doc = fitz.open()
    for path in images:
        if not os.path.isfile(path):
            raise FileNotFoundError(f"文件不存在: {path}")
        img = Image.open(path)
        buf = io.BytesIO()
        img.save(buf, "PNG")
        stream = buf.getvalue()
        if page_size == "auto":
            page = doc.new_page(width=float(img.size[0]), height=float(img.size[1]))
            page.insert_image(page.rect, stream=stream)
        else:
            if page_size == "a4":
                pw, ph = 595.28, 841.89
            elif page_size == "a3":
                pw, ph = 841.89, 1190.55
            elif page_size == "letter":
                pw, ph = 612.0, 792.0
            else:
                raise ValueError(f"不支持的页面尺寸: {page_size}")
            page = doc.new_page(width=pw, height=ph)
            r = fitz.Rect(margin, margin, pw - margin, ph - margin)
            page.insert_image(r, stream=stream, keep_proportion=True)
    if doc.page_count == 0:
        doc.close()
        raise ValueError("没有可转换的图片")
    doc.save(output, garbage=3, deflate=True)
    doc.close()
    return output


def pdf_to_word(input, output, pages=None):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    ensure_parent(output)
    d = Document()
    with _open(input) as doc:
        sel = parse_pages(pages, doc.page_count)
        for idx, p in enumerate(sel):
            page = doc[p]
            if idx > 0:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            blocks = page.get_text("dict", sort=True)
            prev_y = None
            for block in blocks["blocks"]:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    spans = [s for s in line["spans"] if s["text"].strip()]
                    if not spans:
                        continue
                    y = round(line["bbox"][1], 1)
                    para = d.add_paragraph()
                    if prev_y is not None and y - prev_y > 40:
                        para.paragraph_format.space_before = Pt(10)
                    prev_y = y
                    for s in spans:
                        run = para.add_run(s["text"])
                        size = s.get("size") or 11
                        run.font.size = Pt(max(5.0, min(size, 48.0)))
                        flags = s.get("flags") or 0
                        run.font.bold = bool(flags & 16)
                        run.font.italic = bool(flags & 2)
                    if not para.runs:
                        para.add_run("")
    d.save(output)
    return output


def pdf_to_ppt(input, output, mode="image", dpi=150, pages=None):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu

    ensure_parent(output)
    prs = Presentation()
    with _open(input) as doc:
        sel = parse_pages(pages, doc.page_count)
        if not sel:
            raise ValueError("没有可转换的页面")
        first = doc[sel[0]]
        prs.slide_width = Emu(int(first.rect.width * 914400 / 72))
        prs.slide_height = Emu(int(first.rect.height * 914400 / 72))
        blank = prs.slide_layouts[6]
        for p in sel:
            page = doc[p]
            slide = prs.slides.add_slide(blank)
            if mode == "image":
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                buf = io.BytesIO()
                img.save(buf, "PNG")
                buf.seek(0)
                slide.shapes.add_picture(buf, 0, 0, width=prs.slide_width, height=prs.slide_height)
            else:
                text = page.get_text("text", sort=True)
                box = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height)
                tf = box.text_frame
                tf.word_wrap = True
                lines = text.splitlines() or [""]
                tf.text = lines[0]
                for ln in lines[1:]:
                    pgh = tf.add_paragraph()
                    pgh.text = ln
    prs.save(output)
    return output


def pdf_to_excel(input, output, pages=None):
    from openpyxl import Workbook

    ensure_parent(output)
    wb = Workbook()
    wb.remove(wb.active)
    with _open(input) as doc:
        sel = parse_pages(pages, doc.page_count)
        for p in sel:
            page = doc[p]
            ws = wb.create_sheet(title=f"第{p + 1}页")
            placed = False
            try:
                tables = page.find_tables()
            except Exception:
                tables = []
            for t in tables:
                rows = t.extract()
                for r in rows:
                    ws.append(["" if c is None else str(c) for c in r])
                    placed = True
                if rows:
                    ws.append([])
            if not placed:
                for line in page.get_text("text", sort=True).splitlines():
                    ws.append([line])
    if not wb.sheetnames:
        raise ValueError("没有可导出的内容")
    wb.save(output)
    return output
