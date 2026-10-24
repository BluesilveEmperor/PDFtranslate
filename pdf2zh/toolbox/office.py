import os

import pymupdf as fitz

from ._util import ensure_parent

A4_W, A4_H = 595.28, 841.89
A4L_W, A4L_H = 841.89, 595.28
EMU_PT = 72.0 / 914400.0

_CJK_RANGES = (
    (0x2E80, 0x9FFF), (0xA000, 0xA4CF), (0xAC00, 0xD7AF),
    (0xF900, 0xFAFF), (0xFE30, 0xFE4F), (0xFF00, 0xFF60),
)


def _has_cjk(text):
    for ch in text:
        cp = ord(ch)
        for lo, hi in _CJK_RANGES:
            if lo <= cp <= hi:
                return True
    return False


def _font_for(text):
    return "china-s" if _has_cjk(text) else "helv"


def _wrap(text, fontname, size, max_w):
    lines = []
    for seg in text.split("\n"):
        if not seg:
            lines.append("")
            continue
        cur = ""
        last_break = -1
        for ch in seg:
            trial = cur + ch
            if cur and fitz.get_text_length(trial, fontname, size) > max_w:
                if ch == " ":
                    lines.append(cur)
                    cur = ""
                    last_break = -1
                elif last_break >= 0:
                    lines.append(cur[:last_break])
                    cur = cur[last_break + 1:] + ch
                    last_break = -1
                else:
                    lines.append(cur)
                    cur = ch
            else:
                cur = trial
                if ch == " ":
                    last_break = len(cur) - 1
        lines.append(cur)
    return lines


def _write_line(page, x, y, line, size, fontname, bold=False):
    kw = {}
    if bold:
        kw = {"render_mode": 2, "border_width": max(0.2, size * 0.03)}
    page.insert_text((x, y), line, fontsize=size, fontname=fontname, **kw)


def _check_ext(input, exts, name):
    low = input.lower()
    for e in exts:
        if low.endswith(e):
            return
    raise ValueError(f"仅支持 {'/'.join(exts)} 格式（{name} 为私有二进制格式，无法纯 Python 解析）: {input}")


def word_to_pdf(input, output):
    _check_ext(input, [".docx"], ".doc")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if not os.path.isfile(input):
        raise FileNotFoundError(f"文件不存在: {input}")
    ensure_parent(output)
    doc = Document(input)
    pdf = fitz.open()
    margin = 56.0
    max_w = A4_W - 2 * margin
    bottom = A4_H - margin
    page = pdf.new_page(width=A4_W, height=A4_H)
    y = margin

    def new_page():
        nonlocal page, y
        page = pdf.new_page(width=A4_W, height=A4_H)
        y = margin

    def ensure(h):
        nonlocal page, y
        if y + h > bottom:
            new_page()

    def draw_para(para):
        nonlocal page, y
        style = para.style.name if para.style is not None else "Normal"
        if style in ("Heading 1", "标题 1"):
            size, bold = 20.0, True
        elif style in ("Heading 2", "标题 2"):
            size, bold = 16.0, True
        elif style in ("Heading 3", "标题 3"):
            size, bold = 14.0, True
        elif style in ("Title", "标题"):
            size, bold = 24.0, True
        else:
            size, bold = 11.0, False
            if para.style is not None and para.style.font is not None and para.style.font.size is not None:
                size = float(para.style.font.size.pt)
        for run in para.runs:
            if run.bold:
                bold = True
        text = "".join(r.text for r in para.runs) or para.text
        for blob, iw, ih in _iter_inline_images(para, qn):
            ensure(ih)
            if iw > max_w:
                ih = ih * max_w / iw
                iw = max_w
            page.insert_image(fitz.Rect(margin, y, margin + iw, y + ih), stream=blob)
            y += ih + 8
        if not text.strip():
            y += size * 0.6
            return
        fontname = _font_for(text)
        lh = size * 1.4
        align = para.alignment
        for line in _wrap(text, fontname, size, max_w):
            ensure(lh)
            w = fitz.get_text_length(line, fontname, size)
            if align == WD_ALIGN_PARAGRAPH.CENTER:
                x = margin + (max_w - w) / 2
            elif align == WD_ALIGN_PARAGRAPH.RIGHT:
                x = A4_W - margin - w
            else:
                x = margin
            _write_line(page, x, y + size, line, size, fontname, bold)
            y += lh
        y += size * 0.35

    def draw_table(tbl):
        nonlocal page, y
        size = 10.0
        pad = 3.0
        n_cols = len(tbl.columns)
        if n_cols == 0:
            return
        col_w = max_w / n_cols
        for row in tbl.rows:
            cells = [cell.text or "" for cell in row.cells]
            wrapped = [_wrap(c, _font_for(c), size, col_w - 2 * pad) for c in cells]
            n_lines = max((len(w) for w in wrapped), default=1)
            row_h = n_lines * size * 1.35 + 2 * pad
            ensure(row_h)
            for c in range(n_cols):
                rect = fitz.Rect(margin + c * col_w, y, margin + (c + 1) * col_w, y + row_h)
                page.draw_rect(rect, width=0.5)
                fontname = _font_for(cells[c])
                ty = y + pad + size
                for ln in wrapped[c]:
                    page.insert_text((rect.x0 + pad, ty), ln, fontsize=size, fontname=fontname)
                    ty += size * 1.35
            y += row_h
        y += 8

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            draw_para(Paragraph(child, doc))
        elif tag == qn("w:tbl"):
            draw_table(Table(child, doc))
    pdf.save(output, garbage=3, deflate=True)
    pdf.close()
    return output


def _iter_inline_images(para, qn):
    for run in para.runs:
        for blip in run._element.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            try:
                part = para.part.related_parts[rid]
            except Exception:
                continue
            extent = run._element.find(".//" + qn("wp:extent"))
            try:
                w = int(extent.get("cx")) * EMU_PT if extent is not None else 300.0
                h = int(extent.get("cy")) * EMU_PT if extent is not None else 200.0
            except (TypeError, ValueError):
                w, h = 300.0, 200.0
            yield part.blob, w, h


def excel_to_pdf(input, output):
    low = input.lower()
    if not os.path.isfile(input):
        raise FileNotFoundError(f"文件不存在: {input}")
    ensure_parent(output)
    if low.endswith(".xlsx"):
        sheets = _read_xlsx(input)
    elif low.endswith(".xls"):
        sheets = _read_xls(input)
    else:
        raise ValueError("仅支持 .xlsx / .xls 格式（.xlsb/.xlsm 等不受支持）")
    if not sheets:
        raise ValueError("工作簿中没有数据")
    pdf = fitz.open()
    margin = 24.0
    size = 9.0
    line_h = 12.0
    for title, rows in sheets:
        if not rows:
            continue
        n_cols = max(len(r) for r in rows)
        col_w = _col_widths(rows, n_cols, A4L_W - 2 * margin)
        page = pdf.new_page(width=A4L_W, height=A4L_H)
        y = margin
        _write_line(page, margin, y + 12, title, 12, _font_for(title), bold=True)
        y += 24
        for row in rows:
            texts = [str(row[c]) if c < len(row) and row[c] is not None else "" for c in range(n_cols)]
            fonts = [_font_for(t) for t in texts]
            wrapped = [_wrap(t, f, size, col_w[c] - 4) or [""] for c, (t, f) in enumerate(zip(texts, fonts))]
            n_lines = max(len(w) for w in wrapped)
            row_h = n_lines * line_h + 4
            if y + row_h > A4L_H - margin:
                page = pdf.new_page(width=A4L_W, height=A4L_H)
                y = margin
            x = margin
            for c in range(n_cols):
                rect = fitz.Rect(x, y, x + col_w[c], y + row_h)
                page.draw_rect(rect, width=0.4)
                ty = y + size + 2
                for ln in wrapped[c]:
                    page.insert_text((x + 2, ty), ln, fontsize=size, fontname=fonts[c])
                    ty += line_h
                x += col_w[c]
            y += row_h
    if pdf.page_count == 0:
        pdf.close()
        raise ValueError("工作簿中没有数据")
    pdf.save(output, garbage=3, deflate=True)
    pdf.close()
    return output


def _text_width_pt(s, size):
    w = 0.0
    for ch in s:
        cp = ord(ch)
        if any(lo <= cp <= hi for lo, hi in _CJK_RANGES):
            w += size
        else:
            w += size * 0.55
    return w


def _col_widths(rows, n_cols, usable_w):
    widths = [30.0] * n_cols
    for row in rows:
        for c in range(min(len(row), n_cols)):
            v = row[c]
            if v is None:
                continue
            s = str(v)
            widths[c] = max(widths[c], _text_width_pt(s, 9.0) + 6)
    total = sum(widths)
    if total > usable_w:
        scale = usable_w / total
        widths = [max(24.0, w * scale) for w in widths]
        total = sum(widths)
        if total > usable_w:
            scale = usable_w / total
            widths = [w * scale for w in widths]
    return widths


def _read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else v for v in row]
            if any(str(v).strip() for v in vals):
                rows.append(vals)
            if len(rows) >= 5000:
                break
        sheets.append((ws.title or "Sheet", rows))
    wb.close()
    return sheets


def _read_xls(path):
    import xlrd

    book = xlrd.open_workbook(path)
    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(min(sh.nrows, 5000)):
            vals = [sh.cell_value(r, c) for c in range(sh.ncols)]
            if any(str(v).strip() for v in vals):
                rows.append(vals)
        sheets.append((sh.name or "Sheet", rows))
    return sheets


def ppt_to_pdf(input, output):
    _check_ext(input, [".pptx"], ".ppt")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if not os.path.isfile(input):
        raise FileNotFoundError(f"文件不存在: {input}")
    ensure_parent(output)
    prs = Presentation(input)
    pdf = fitz.open()
    slide_w = prs.slide_width * EMU_PT
    slide_h = prs.slide_height * EMU_PT
    if slide_w <= 0 or slide_h <= 0:
        slide_w, slide_h = 960.0, 540.0
    for slide in prs.slides:
        page = pdf.new_page(width=slide_w, height=slide_h)
        for shape in slide.shapes:
            try:
                left = shape.left * EMU_PT
                top = shape.top * EMU_PT
                w = shape.width * EMU_PT
                h = shape.height * EMU_PT
            except (TypeError, AttributeError):
                continue
            if w <= 0 or h <= 0:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    page.insert_image(fitz.Rect(left, top, left + w, top + h), stream=shape.image.blob)
                except Exception:
                    pass
                continue
            if not shape.has_text_frame:
                continue
            ty = top
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs) or para.text
                if not text.strip():
                    ty += 10
                    continue
                size = 14.0
                for run in para.runs:
                    if run.font.size is not None:
                        size = float(run.font.size.pt)
                        break
                fontname = _font_for(text)
                lh = size * 1.25
                for line in _wrap(text, fontname, size, w):
                    _write_line(page, left, ty + size, line, size, fontname)
                    ty += lh
                ty += 4
    if pdf.page_count == 0:
        pdf.close()
        raise ValueError("演示文稿中没有幻灯片")
    pdf.save(output, garbage=3, deflate=True)
    pdf.close()
    return output
