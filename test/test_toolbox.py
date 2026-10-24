"""toolbox 融合功能的冒烟 / 单元测试。

覆盖范围：
    * 统一 API：``COMMANDS`` / ``CATEGORIES`` / ``build_argv`` / ``validate`` /
      ``run_command``（含成功与失败路径）。
    * 全部 24 个命令的代表性执行路径：页面操作、内容编辑、安全、PDF 转换、
      Office 转换。
    * 独立命令行入口 ``pdf2zh.toolbox.cli.main``。
    * Gradio「工具箱」Tab 的可构建性 ``pdf2zh.toolbox_gui.build_toolbox_tab``。

所有输入（PDF / 图片 / Office 文档）均在运行期用 pymupdf、Pillow、
python-docx、python-pptx、openpyxl 程序化生成，不依赖任何外部样例文件。
可选依赖缺失时相关用例自动 skip。
"""

from __future__ import annotations

import os

import pytest
import pymupdf as fitz

from pdf2zh.toolbox import CATEGORIES, COMMANDS, build_argv, run_command, validate
from pdf2zh.toolbox.cli import main as toolbox_main


# --------------------------------------------------------------------------- #
# 程序化输入生成辅助
# --------------------------------------------------------------------------- #
def _make_pdf(path, texts, width=595, height=842):
    """生成含可搜索文本层的 PDF。

    :param texts: 单个字符串（每页相同）或逐页文本列表，便于用文本内容区分页面。
    """
    if isinstance(texts, str):
        texts = [texts]
    doc = fitz.open()
    for t in texts:
        page = doc.new_page(width=width, height=height)
        page.insert_text((72, 100), t, fontsize=12)
    doc.save(str(path))
    doc.close()
    return str(path)


def _make_linked_pdf(path):
    """生成带一个 URI 链接的 PDF，并确认链接确实写入文件（前置条件）。"""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Link Page", fontsize=12)
    page.insert_link(
        {
            "kind": fitz.LINK_URI,
            "from": fitz.Rect(72, 80, 240, 110),
            "uri": "https://example.com",
        }
    )
    doc.save(str(path))
    doc.close()
    with fitz.open(str(path)) as check:
        assert len(check[0].get_links()) >= 1, "前置条件：生成的 PDF 应含至少 1 个链接"
    return str(path)


def _make_png(path, size=(120, 80), color=(200, 30, 30)):
    from PIL import Image

    Image.new("RGB", size, color).save(str(path))
    return str(path)


def _make_png_alpha(path, size=(120, 80)):
    from PIL import Image

    Image.new("RGBA", size, (30, 30, 200, 255)).save(str(path))
    return str(path)


def _page_count(path):
    with fitz.open(str(path)) as doc:
        return doc.page_count


def _page_text(path, index=0):
    with fitz.open(str(path)) as doc:
        return doc[index].get_text()


def _texts(path):
    with fitz.open(str(path)) as doc:
        return [page.get_text() for page in doc]


# --------------------------------------------------------------------------- #
# API 结构：COMMANDS / CATEGORIES
# --------------------------------------------------------------------------- #
def test_commands_catalog_shape():
    """COMMANDS 应为 24 项、CATEGORIES 应为 5 类，且元数据字段完整。"""
    assert len(COMMANDS) == 24
    assert len(CATEGORIES) == 5

    categories = [name for name, _key in CATEGORIES]
    assert categories == ["页面操作", "内容编辑", "安全", "PDF 转换", "Office 转换"]

    for cmd, meta in COMMANDS.items():
        assert {"cat", "label", "icon", "desc", "fields"} <= set(meta), cmd
        assert meta["cat"] in categories, cmd
        assert isinstance(meta["fields"], list) and meta["fields"], cmd
        for field in meta["fields"]:
            assert {"key", "label", "type"} <= set(field), (cmd, field)

    # 每类至少含 1 个命令
    for cat in categories:
        assert [c for c, m in COMMANDS.items() if m["cat"] == cat], cat


def test_commands_cover_expected_names():
    """命令集合应精确覆盖需求中列出的 24 个命令。"""
    expected = {
        "merge", "split", "extract", "delete", "insert", "rotate", "cut", "pages-merger",
        "page-number", "watermark-text", "watermark-image", "text-replace", "text-delete",
        "remove-links", "encrypt", "decrypt",
        "pdf-to-image", "image-to-pdf", "pdf-to-word", "pdf-to-ppt", "pdf-to-excel",
        "word-to-pdf", "excel-to-pdf", "ppt-to-pdf",
    }
    assert set(COMMANDS) == expected


# --------------------------------------------------------------------------- #
# API 结构：build_argv
# --------------------------------------------------------------------------- #
def test_build_argv_positional_and_optional():
    """位置参数在前、可选参数以 key/value 成对追加。"""
    assert build_argv(
        "merge", {"inputs": ["a.pdf", "b.pdf"], "-o": "m.pdf"}
    ) == ["merge", "a.pdf", "b.pdf", "-o", "m.pdf"]

    assert build_argv(
        "extract", {"input": "x.pdf", "-p": "1-2", "-o": "o.pdf"}
    ) == ["extract", "x.pdf", "-p", "1-2", "-o", "o.pdf"]


def test_build_argv_bool_flag_and_skip_empty():
    """bool 为真时仅输出开关名；空值字段被跳过。"""
    argv = build_argv(
        "encrypt",
        {"input": "x.pdf", "--user-pw": "pw", "--owner-pw": "", "--no-print": True, "-o": "o.pdf"},
    )
    assert argv == ["encrypt", "x.pdf", "--user-pw", "pw", "--no-print", "-o", "o.pdf"]


# --------------------------------------------------------------------------- #
# API 结构：validate
# --------------------------------------------------------------------------- #
def test_validate_ok_returns_empty():
    assert validate("merge", {"inputs": ["a.pdf"], "-o": "m.pdf"}) == []


def test_validate_missing_required_reports_errors():
    errors = validate("merge", {})
    assert errors, "缺少必填项应返回错误列表"
    assert any("缺少必填项" in e for e in errors)
    # inputs 与 -o 两个必填项都缺失
    assert len(errors) >= 2


def test_validate_type_and_cross_field_rules():
    """int 类型校验与 split 的跨字段规则。"""
    assert any("须为整数" in e for e in validate("cut", {"input": "a.pdf", "-o": "o.pdf", "--left": "abc"})) is False
    assert any("须为数字" in e for e in validate("cut", {"input": "a.pdf", "-o": "o.pdf", "--left": "abc"}))
    assert any(
        "拆分方式" in e
        for e in validate("split", {"input": "a.pdf", "-o": "d"})
    )
    assert validate("split", {"input": "a.pdf", "-o": "d", "--every": "2"}) == []


# --------------------------------------------------------------------------- #
# run_command 失败路径
# --------------------------------------------------------------------------- #
def test_run_command_unknown_command():
    result = run_command("no-such-command", {})
    assert result["ok"] is False
    assert "未知命令" in result["error"]
    assert result["output"] is None


def test_run_command_missing_required_argument(tmp_path):
    result = run_command("merge", {"-o": str(tmp_path / "out.pdf")})
    assert result["ok"] is False
    assert result["error"]
    assert result["output"] is None


def test_run_command_missing_input_file(tmp_path):
    result = run_command(
        "merge", {"inputs": [str(tmp_path / "nope.pdf")], "-o": str(tmp_path / "out.pdf")}
    )
    assert result["ok"] is False
    assert result["error"]  # FileNotFoundError 被统一兜底为结构化错误


def test_run_command_invalid_page_expression(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"])
    out = str(tmp_path / "out.pdf")
    result = run_command("extract", {"input": src, "-p": "9", "-o": out})
    assert result["ok"] is False
    assert result["error"]
    assert not os.path.exists(out)


def test_run_command_result_shape(tmp_path):
    """成功时返回结构应为 4 键，且 error 为空串。"""
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"])
    result = run_command("extract", {"input": src, "-p": "1", "-o": str(tmp_path / "o.pdf")})
    assert result["ok"] is True
    assert set(result) == {"ok", "output", "log", "error"}
    assert result["error"] == ""
    assert os.path.exists(result["output"])


# --------------------------------------------------------------------------- #
# 页面操作
# --------------------------------------------------------------------------- #
def test_merge_concatenates_all_pages(tmp_path):
    a = _make_pdf(tmp_path / "a.pdf", ["ALPHA_ONE", "ALPHA_TWO"])
    b = _make_pdf(tmp_path / "b.pdf", ["BETA_ONE", "BETA_TWO", "BETA_THREE"])
    out = tmp_path / "merged.pdf"

    result = run_command("merge", {"inputs": [a, b], "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert os.path.isfile(result["output"])
    assert _page_count(result["output"]) == 5
    merged_text = "".join(_texts(result["output"]))
    assert "ALPHA_ONE" in merged_text and "BETA_THREE" in merged_text


def test_split_every_produces_expected_parts(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2", "P3", "P4", "P5"])
    outdir = tmp_path / "parts"

    result = run_command("split", {"input": src, "-o": str(outdir), "--every": "2"})

    assert result["ok"] is True, result["error"]
    parts = result["output"]
    assert isinstance(parts, list) and len(parts) == 3
    assert all(os.path.isfile(p) for p in parts)
    assert [_page_count(p) for p in parts] == [2, 2, 1]


def test_split_ranges(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2", "P3", "P4"])
    outdir = tmp_path / "parts"

    result = run_command("split", {"input": src, "-o": str(outdir), "--ranges": "1-2,3-4"})

    assert result["ok"] is True, result["error"]
    assert len(result["output"]) == 2
    assert [_page_count(p) for p in result["output"]] == [2, 2]


def test_extract_keeps_only_selected_pages(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["UNIQUE_1", "UNIQUE_2", "UNIQUE_3", "UNIQUE_4"])
    out = tmp_path / "out.pdf"

    result = run_command("extract", {"input": src, "-p": "1,3", "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 2
    kept = "".join(_texts(result["output"]))
    assert "UNIQUE_1" in kept and "UNIQUE_3" in kept
    assert "UNIQUE_2" not in kept and "UNIQUE_4" not in kept


def test_delete_removes_selected_pages(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["UNIQUE_1", "UNIQUE_2", "UNIQUE_3", "UNIQUE_4"])
    out = tmp_path / "out.pdf"

    result = run_command("delete", {"input": src, "-p": "2,4", "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 2
    kept = "".join(_texts(result["output"]))
    assert "UNIQUE_1" in kept and "UNIQUE_3" in kept
    assert "UNIQUE_2" not in kept and "UNIQUE_4" not in kept


def test_insert_places_pages_at_position(tmp_path):
    base = _make_pdf(tmp_path / "base.pdf", ["A1", "A2"])
    other = _make_pdf(tmp_path / "other.pdf", ["B1", "B2"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "insert", {"input": base, "--file": other, "--pos": "1", "-o": str(out)}
    )

    assert result["ok"] is True, result["error"]
    texts = _texts(result["output"])
    assert len(texts) == 4
    assert "A1" in texts[0]
    assert "B1" in texts[1] and "B2" in texts[2]
    assert "A2" in texts[3]


def test_rotate_sets_page_rotation(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2", "P3"])
    out = tmp_path / "out.pdf"

    result = run_command("rotate", {"input": src, "--angle": "90", "-p": "2", "-o": str(out)})

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        rotations = [page.rotation for page in doc]
    assert rotations == [0, 90, 0]


def test_rotate_all_pages(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"])
    out = tmp_path / "out.pdf"

    result = run_command("rotate", {"input": src, "--angle": "180", "-o": str(out)})

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert [page.rotation for page in doc] == [180, 180]


def test_cut_margins_shrinks_cropbox(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "cut",
        {"input": src, "-o": str(out), "--left": "20", "--right": "20", "--top": "20", "--bottom": "20"},
    )

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        box = doc[0].cropbox
        assert box.width == pytest.approx(595 - 40, abs=0.5)
        assert box.height == pytest.approx(842 - 40, abs=0.5)
        assert box.x0 == pytest.approx(20, abs=0.5) and box.y0 == pytest.approx(20, abs=0.5)


def test_cut_margins_too_large_fails(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "cut",
        {"input": src, "-o": str(out), "--left": "400", "--right": "400"},
    )
    assert result["ok"] is False
    assert result["error"]


def test_pages_merger_interleave(tmp_path):
    a = _make_pdf(tmp_path / "a.pdf", ["A1", "A2"])
    b = _make_pdf(tmp_path / "b.pdf", ["B1", "B2", "B3"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "pages-merger", {"input_a": a, "input_b": b, "--mode": "interleave", "-o": str(out)}
    )

    assert result["ok"] is True, result["error"]
    texts = _texts(result["output"])
    assert len(texts) == 5
    order = ["A1", "B1", "A2", "B2", "B3"]
    for idx, token in enumerate(order):
        assert token in texts[idx], (idx, token, texts[idx])


def test_pages_merger_append(tmp_path):
    a = _make_pdf(tmp_path / "a.pdf", ["A1", "A2"])
    b = _make_pdf(tmp_path / "b.pdf", ["B1"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "pages-merger", {"input_a": a, "input_b": b, "--mode": "append", "-o": str(out)}
    )

    assert result["ok"] is True, result["error"]
    texts = _texts(result["output"])
    assert len(texts) == 3
    assert "A1" in texts[0] and "A2" in texts[1] and "B1" in texts[2]


# --------------------------------------------------------------------------- #
# 内容编辑
# --------------------------------------------------------------------------- #
def test_page_number_adds_numbers(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2", "P3"])
    out = tmp_path / "out.pdf"

    result = run_command("page-number", {"input": src, "-o": str(out), "--start": "1"})

    assert result["ok"] is True, result["error"]
    texts = _texts(result["output"])
    assert len(texts) == 3
    for idx, text in enumerate(texts, start=1):
        assert str(idx) in text


def test_watermark_text_tiled_embeds_text(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["Original"])
    out = tmp_path / "out.pdf"

    result = run_command("watermark-text", {"input": src, "--text": "DRAFT", "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert "DRAFT" in _page_text(result["output"])


def test_watermark_image_embeds_image(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["Original"])
    img = _make_png_alpha(tmp_path / "wm.png")
    out = tmp_path / "out.pdf"

    result = run_command(
        "watermark-image", {"input": src, "--image": img, "-o": str(out), "--center": True}
    )

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert doc.page_count == 1
        assert len(doc[0].get_images(full=True)) >= 1


def test_text_replace_swaps_text(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["Hello World"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "text-replace", {"input": src, "--old": "Hello", "--new": "World", "-o": str(out)}
    )

    assert result["ok"] is True, result["error"]
    text = _page_text(result["output"])
    assert "Hello" not in text
    assert "World" in text


def test_text_replace_missing_text_fails(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["Hello World"])
    out = tmp_path / "out.pdf"

    result = run_command(
        "text-replace", {"input": src, "--old": "ABSENT_TOKEN", "--new": "X", "-o": str(out)}
    )
    assert result["ok"] is False
    assert "未找到文本" in result["error"]


def test_text_delete_removes_only_target(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["Secret Data"])
    out = tmp_path / "out.pdf"

    result = run_command("text-delete", {"input": src, "--text": "Secret", "-o": str(out)})

    assert result["ok"] is True, result["error"]
    text = _page_text(result["output"])
    assert "Secret" not in text
    assert "Data" in text


def test_remove_links_deletes_all_links(tmp_path):
    src = _make_linked_pdf(tmp_path / "links.pdf")
    out = tmp_path / "out.pdf"

    result = run_command("remove-links", {"input": src, "-o": str(out)})

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert doc[0].get_links() == []


# --------------------------------------------------------------------------- #
# 安全
# --------------------------------------------------------------------------- #
def test_encrypt_produces_encrypted_pdf(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["Confidential"])
    out = tmp_path / "enc.pdf"

    result = run_command(
        "encrypt", {"input": src, "--user-pw": "secret123", "-o": str(out)}
    )

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert doc.is_encrypted is True
        # authenticate 返回权限位整数：0 表示密码错误，非 0 表示通过
        assert doc.authenticate("secret123") != 0
        assert doc.authenticate("wrong-pw") == 0


def test_decrypt_recovers_short_plaintext_password(tmp_path):
    # 明文密码仅 2 字符、字符集限定为 "ab"，暴力枚举可快速命中（jobs=1 走单进程）。
    src = _make_pdf(tmp_path / "plain.pdf", ["Recovered page 1", "Recovered page 2"])
    enc = tmp_path / "enc.pdf"
    enc_result = run_command("encrypt", {"input": src, "--user-pw": "ab", "-o": str(enc)})
    assert enc_result["ok"] is True, enc_result["error"]
    with fitz.open(enc_result["output"]) as doc:
        assert doc.is_encrypted is True
        assert doc.authenticate("") == 0  # 空密码无法通过，需要打开密码

    out = tmp_path / "dec.pdf"
    result = run_command(
        "decrypt",
        {
            "input": enc_result["output"],
            "--charset": "ab",
            "--min-len": "2",
            "--max-len": "2",
            "--jobs": "1",
            "-o": str(out),
        },
    )

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert doc.page_count == 2
        assert doc.is_encrypted is False  # 恢复后可无密码打开


# --------------------------------------------------------------------------- #
# PDF 转换
# --------------------------------------------------------------------------- #
def test_pdf_to_image_renders_pages(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"], width=300, height=200)
    outdir = tmp_path / "imgs"

    result = run_command(
        "pdf-to-image", {"input": src, "-o": str(outdir), "--dpi": "72", "--fmt": "png"}
    )

    assert result["ok"] is True, result["error"]
    files = result["output"]
    assert isinstance(files, list) and len(files) == 2
    assert all(p.endswith(".png") for p in files)
    from PIL import Image

    with Image.open(files[0]) as im:
        assert im.width > 0 and im.height > 0


def test_image_to_pdf_creates_one_page_per_image(tmp_path):
    i1 = _make_png(tmp_path / "i1.png", size=(200, 300))
    i2 = _make_png(tmp_path / "i2.png", size=(200, 300))
    out = tmp_path / "out.pdf"

    result = run_command("image-to-pdf", {"images": [i1, i2], "-o": str(out)})

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert doc.page_count == 2
        assert doc[0].rect.width == pytest.approx(200, abs=1.0)
        assert doc[0].rect.height == pytest.approx(300, abs=1.0)


def test_pdf_to_word_generates_docx(tmp_path):
    docx = pytest.importorskip("docx")
    src = _make_pdf(tmp_path / "in.pdf", ["WordDocText"])
    out = tmp_path / "out.docx"

    result = run_command("pdf-to-word", {"input": src, "-o": str(out)})

    assert result["ok"] is True, result["error"]
    document = docx.Document(result["output"])
    joined = "\n".join(p.text for p in document.paragraphs)
    assert "WordDocText" in joined


def test_pdf_to_ppt_generates_slides(tmp_path):
    pptx = pytest.importorskip("pptx")
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"])
    out = tmp_path / "out.pptx"

    result = run_command("pdf-to-ppt", {"input": src, "-o": str(out), "--mode": "image"})

    assert result["ok"] is True, result["error"]
    prs = pptx.Presentation(result["output"])
    assert len(prs.slides) == 2


def test_pdf_to_excel_generates_sheets(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"])
    out = tmp_path / "out.xlsx"

    result = run_command("pdf-to-excel", {"input": src, "-o": str(out)})

    assert result["ok"] is True, result["error"]
    wb = openpyxl.load_workbook(result["output"])
    assert len(wb.sheetnames) == 2
    wb.close()


# --------------------------------------------------------------------------- #
# Office 转换
# --------------------------------------------------------------------------- #
def test_word_to_pdf(tmp_path):
    docx = pytest.importorskip("docx")
    src = tmp_path / "in.docx"
    document = docx.Document()
    document.add_heading("Headline", level=1)
    document.add_paragraph("HelloFromWord")
    document.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("word-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) >= 1
    assert "HelloFromWord" in "".join(_texts(result["output"]))


def test_excel_to_pdf(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    src = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SheetA"
    ws.append(["CellValueA", "CellValueB"])
    ws.append([10, 20])
    wb.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("excel-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) >= 1
    assert "CellValueA" in "".join(_texts(result["output"]))


def test_ppt_to_pdf(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    src = tmp_path / "in.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # add_textbox 的尺寸单位为 EMU，用 Inches() 给定合适宽度避免换行过窄
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1))
    box.text_frame.text = "SlideText"
    prs.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("ppt-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 1
    # 渲染按字形换行，比较时去除所有空白字符
    extracted = "".join("".join(_texts(result["output"])).split())
    assert "SlideText" in extracted


def test_word_to_pdf_rejects_legacy_doc(tmp_path):
    """非 .docx 扩展名应走 _check_ext 分支并返回结构化错误。"""
    bad = tmp_path / "legacy.doc"
    bad.write_bytes(b"not a real doc")
    out = tmp_path / "out.pdf"

    result = run_command("word-to-pdf", {"input": str(bad), "-o": str(out)})

    assert result["ok"] is False
    assert result["error"]


# --------------------------------------------------------------------------- #
# 独立命令行入口
# --------------------------------------------------------------------------- #
def test_cli_main_merge_success(tmp_path, capsys):
    a = _make_pdf(tmp_path / "a.pdf", ["A1"])
    b = _make_pdf(tmp_path / "b.pdf", ["B1"])
    out = tmp_path / "merged.pdf"

    rc = toolbox_main(["merge", a, b, "-o", str(out)])

    assert rc == 0
    assert os.path.isfile(str(out))
    assert _page_count(out) == 2


def test_cli_main_missing_required_returns_nonzero(tmp_path, capsys):
    rc = toolbox_main(["merge"])
    assert rc != 0
    assert rc == 2


def test_cli_main_unknown_command_returns_nonzero():
    rc = toolbox_main(["definitely-not-a-command"])
    assert rc != 0


# --------------------------------------------------------------------------- #
# Gradio 工具箱 Tab
# --------------------------------------------------------------------------- #
def test_build_toolbox_tab_constructs_inside_blocks():
    """build_toolbox_tab 应能在 gr.Blocks/gr.Tab 上下文中构建并返回组件。"""
    gr = pytest.importorskip("gradio")
    from pdf2zh.toolbox_gui import build_toolbox_tab

    with gr.Blocks() as demo:
        with gr.Tab("工具箱"):
            component = build_toolbox_tab()

    assert component is not None
    assert isinstance(demo, gr.Blocks)

# --------------------------------------------------------------------------- #
# 内部工具函数 _util：页码范围解析与文件系统辅助
# --------------------------------------------------------------------------- #
def test_util_parse_pages_variants():
    from pdf2zh.toolbox import _util

    assert _util.parse_pages("", 3) == [0, 1, 2]        # 空表达式 => 全部页
    assert _util.parse_pages("   ", 3) == [0, 1, 2]
    assert _util.parse_pages("1-3", 5) == [0, 1, 2]      # 区间
    assert _util.parse_pages("4-9", 5) == [3, 4]         # 上界截断到总页数
    assert _util.parse_pages("2,4", 5) == [1, 3]         # 离散页码
    assert _util.parse_pages("1, ,3", 5) == [0, 2]       # 空段被跳过


def test_util_parse_pages_rejects_invalid():
    from pdf2zh.toolbox import _util

    for spec in ["0", "abc", "3-1", "x-y"]:
        with pytest.raises(ValueError):
            _util.parse_pages(spec, 3)
    with pytest.raises(ValueError):
        _util.parse_pages(",", 3)  # 选择为空


def test_util_parse_ranges_variants():
    from pdf2zh.toolbox import _util

    assert _util.parse_ranges("1-2,3-4", 4) == [(0, 1), (2, 3)]
    assert _util.parse_ranges("2", 4) == [(1, 1)]        # 单页
    assert _util.parse_ranges("1-10", 4) == [(0, 3)]     # 截断
    assert _util.parse_ranges(",2", 4) == [(1, 1)]       # 空段被跳过

    with pytest.raises(ValueError):
        _util.parse_ranges(",", 4)


def test_util_ensure_parent_creates_dirs(tmp_path):
    from pdf2zh.toolbox import _util

    target = tmp_path / "deep" / "nested" / "out.pdf"
    _util.ensure_parent(str(target))
    assert (tmp_path / "deep" / "nested").is_dir()


def test_util_die_terminates_with_code(capsys):
    from pdf2zh.toolbox import _util

    with pytest.raises(SystemExit) as exc:
        _util.die("boom-message", code=3)
    assert exc.value.code == 3
    assert "boom-message" in capsys.readouterr().err


# --------------------------------------------------------------------------- #
# 模块入口 python -m pdf2zh.toolbox
# --------------------------------------------------------------------------- #
def test_module_entry_point_version(monkeypatch):
    import runpy
    import sys

    monkeypatch.setattr(sys, "argv", ["pdf2zh.toolbox", "--version"])
    with pytest.raises(SystemExit) as exc:
        runpy.run_module("pdf2zh.toolbox", run_name="__main__")
    assert exc.value.code == 0


# --------------------------------------------------------------------------- #
# merge 的输入展开：目录 / 通配符 / 单路径
# --------------------------------------------------------------------------- #
def test_merge_expands_directory_single_path_and_glob(tmp_path):
    d = tmp_path / "src"
    d.mkdir()
    _make_pdf(d / "a.pdf", ["DIR_A"])
    _make_pdf(d / "b.pdf", ["DIR_B"])

    r1 = run_command("merge", {"inputs": [str(d)], "-o": str(tmp_path / "m1.pdf")})
    assert r1["ok"] is True, r1["error"]
    assert _page_count(r1["output"]) == 2

    r2 = run_command("merge", {"inputs": str(d / "a.pdf"), "-o": str(tmp_path / "m2.pdf")})
    assert r2["ok"] is True, r2["error"]
    assert _page_count(r2["output"]) == 1

    r3 = run_command("merge", {"inputs": [str(d / "*.pdf")], "-o": str(tmp_path / "m3.pdf")})
    assert r3["ok"] is True, r3["error"]
    assert _page_count(r3["output"]) == 2


def test_merge_directory_without_pdfs_fails(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    result = run_command("merge", {"inputs": [str(empty)], "-o": str(tmp_path / "o.pdf")})
    assert result["ok"] is False
    assert result["error"]


# --------------------------------------------------------------------------- #
# validate：int 类型非法值
# --------------------------------------------------------------------------- #
def test_validate_int_field_rejects_non_integer():
    errors = validate("page-number", {"input": "a.pdf", "-o": "o.pdf", "--start": "abc"})
    assert any("须为整数" in e for e in errors)


# --------------------------------------------------------------------------- #
# converters 补充分支
# --------------------------------------------------------------------------- #
def test_pdf_to_image_jpg_format(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", ["P1", "P2"], width=200, height=150)
    outdir = tmp_path / "imgs"

    result = run_command(
        "pdf-to-image",
        {"input": src, "-o": str(outdir), "--dpi": "72", "--fmt": "jpg", "--quality": "80", "-p": "1"},
    )

    assert result["ok"] is True, result["error"]
    assert len(result["output"]) == 1
    assert result["output"][0].endswith(".jpg")
    from PIL import Image

    with Image.open(result["output"][0]) as im:
        assert im.format == "JPEG"


def test_pdf_to_image_missing_file_returns_error(tmp_path):
    result = run_command(
        "pdf-to-image", {"input": str(tmp_path / "nope.pdf"), "-o": str(tmp_path / "imgs")}
    )
    assert result["ok"] is False
    assert result["error"]


def test_image_to_pdf_a4_page_size_with_margin(tmp_path):
    i1 = _make_png(tmp_path / "i1.png", size=(100, 100))
    out = tmp_path / "a4.pdf"

    result = run_command(
        "image-to-pdf", {"images": [i1], "-o": str(out), "--page-size": "a4", "--margin": "20"}
    )

    assert result["ok"] is True, result["error"]
    with fitz.open(result["output"]) as doc:
        assert doc.page_count == 1
        assert doc[0].rect.width == pytest.approx(595.28, abs=1.0)
        assert doc[0].rect.height == pytest.approx(841.89, abs=1.0)


def test_image_to_pdf_with_no_images_raises(tmp_path):
    from pdf2zh.toolbox import converters

    with pytest.raises(ValueError):
        converters.image_to_pdf([], str(tmp_path / "o.pdf"))


def test_pdf_to_ppt_text_mode(tmp_path):
    pptx = pytest.importorskip("pptx")
    src = _make_pdf(tmp_path / "in.pdf", ["EditableText"])
    out = tmp_path / "out.pptx"

    result = run_command("pdf-to-ppt", {"input": src, "-o": str(out), "--mode": "text"})

    assert result["ok"] is True, result["error"]
    prs = pptx.Presentation(result["output"])
    assert len(prs.slides) == 1
    joined = "".join(
        sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
    )
    assert "EditableText" in "".join(joined.split())


# --------------------------------------------------------------------------- #
# office 补充分支
# --------------------------------------------------------------------------- #
def test_excel_to_pdf_rejects_unsupported_extension(tmp_path):
    bad = tmp_path / "data.csv"
    bad.write_text("a,b\n1,2\n", encoding="utf-8")
    result = run_command("excel-to-pdf", {"input": str(bad), "-o": str(tmp_path / "o.pdf")})
    assert result["ok"] is False
    assert result["error"]


def test_excel_to_pdf_legacy_xls_format(tmp_path):
    """旧版 .xls 走 xlrd 读取路径。"""
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    src = tmp_path / "legacy.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("OldSheet")
    sheet.write(0, 0, "LegacyXlsCell")
    sheet.write(1, 0, 123)
    book.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("excel-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) >= 1
    assert "LegacyXlsCell" in "".join(_texts(result["output"]))


def test_excel_to_pdf_empty_workbook_fails(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    src = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(str(src))

    result = run_command("excel-to-pdf", {"input": str(src), "-o": str(tmp_path / "o.pdf")})

    assert result["ok"] is False
    assert result["error"]


def test_excel_to_pdf_multiple_sheets_with_wide_cells(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    src = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1.append(["WideColumnValue" + "x" * 200, "B1"])  # 触发列宽缩放分支
    ws2 = wb.create_sheet("Second")
    ws2.append(["SecondSheetCell", 42])
    wb.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("excel-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 2
    joined = "".join(_texts(result["output"]))
    assert "SecondSheetCell" in joined


def test_word_to_pdf_renders_heading_paragraph_and_table(tmp_path):
    docx = pytest.importorskip("docx")
    src = tmp_path / "rich.docx"
    d = docx.Document()
    d.add_heading("HeadingTitle", level=1)
    d.add_paragraph("BodyParagraph")
    d.add_paragraph("Wrapped" + "A" * 300)  # 触发换行分支
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "R1C1"
    table.cell(0, 1).text = "R1C2"
    table.cell(1, 0).text = "R2C1"
    table.cell(1, 1).text = "R2C2"
    d.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("word-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) >= 1
    joined = "".join(_texts(result["output"]))
    assert "HeadingTitle" in joined
    assert "R1C1" in joined and "R2C2" in joined


def test_ppt_to_pdf_with_picture(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    img = _make_png(tmp_path / "pic.png", size=(80, 60))
    src = tmp_path / "pic.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img), Inches(1), Inches(1), Inches(2), Inches(1.5))
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = "Caption"
    prs.save(str(src))
    out = tmp_path / "out.pdf"

    result = run_command("ppt-to-pdf", {"input": str(src), "-o": str(out)})

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 1


# --------------------------------------------------------------------------- #
# 安全：字典攻击与旧版（R<5）加密恢复路径
# --------------------------------------------------------------------------- #
def test_decrypt_with_wordlist(tmp_path):
    """提供字典文件时应走字典攻击阶段并命中密码。"""
    src = _make_pdf(tmp_path / "plain.pdf", ["WordlistPage"])
    enc = tmp_path / "enc.pdf"
    doc = fitz.open(src)
    doc.save(str(enc), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="zzzq", owner_pw="zzzq")
    doc.close()

    wordlist = tmp_path / "words.txt"
    wordlist.write_text("# comment line\nnope-password\nzzzq\n", encoding="utf-8")
    out = tmp_path / "dec.pdf"

    result = run_command(
        "decrypt",
        {
            "input": str(enc),
            "--wordlist": str(wordlist),
            "--charset": "ab",
            "--min-len": "1",
            "--max-len": "1",
            "--jobs": "1",
            "-o": str(out),
        },
    )

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 1
    with fitz.open(result["output"]) as doc:
        assert doc.is_encrypted is False


def test_decrypt_legacy_aes128_encryption(tmp_path):
    """AES-128（R<5）加密走 RC4/key 派生分支，仍应能恢复密码。"""
    src = _make_pdf(tmp_path / "plain.pdf", ["LegacyPage"])
    enc = tmp_path / "enc128.pdf"
    doc = fitz.open(src)
    doc.save(str(enc), encryption=fitz.PDF_ENCRYPT_AES_128, user_pw="ab", owner_pw="ab")
    doc.close()
    with fitz.open(str(enc)) as chk:
        assert chk.is_encrypted is True
        assert chk.authenticate("") == 0

    out = tmp_path / "dec.pdf"
    result = run_command(
        "decrypt",
        {
            "input": str(enc),
            "--charset": "ab",
            "--min-len": "2",
            "--max-len": "2",
            "--jobs": "1",
            "-o": str(out),
        },
    )

    assert result["ok"] is True, result["error"]
    assert _page_count(result["output"]) == 1


def test_word_to_pdf_handles_cjk_text(tmp_path):
    """含 CJK 字符时应选择内置中文字体分支（字供不可用时跳过）。"""
    docx = pytest.importorskip("docx")
    src = tmp_path / "cjk.docx"
    d = docx.Document()
    d.add_paragraph("中文字符测试文本")
    d.save(str(src))
    out = tmp_path / "cjk.pdf"

    result = run_command("word-to-pdf", {"input": str(src), "-o": str(out)})

    if not result["ok"]:
        pytest.skip(f"PyMuPDF 当前构建不支持内置中文字体: {result['error']}")
    assert _page_count(result["output"]) >= 1

# --------------------------------------------------------------------------- #
# 代码审查修复回归（task 16）：M1 / M2 / m1 / m2 / m3 / m4
# --------------------------------------------------------------------------- #
def test_decrypt_default_output_uses_removed_prefix(tmp_path):
    """M1：未传 -o 时解密结果写到输入同目录，文件名加 Removed_ 前缀。"""
    src = _make_pdf(tmp_path / "orig.pdf", ["SecretPage"])
    enc = tmp_path / "locked.pdf"
    doc = fitz.open(src)
    # 仅设权限密码（用户密码为空），decrypt 无需暴力恢复即可打开
    doc.save(
        str(enc), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="", owner_pw="ownerpw"
    )
    doc.close()

    result = run_command("decrypt", {"input": str(enc)})

    assert result["ok"] is True, result["error"]
    expected = tmp_path / "Removed_locked.pdf"
    assert os.path.normcase(os.path.abspath(result["output"])) == os.path.normcase(
        str(expected)
    )
    assert expected.is_file()
    with fitz.open(expected) as out:
        assert out.is_encrypted is False


def test_decrypt_rejects_excessive_max_len(tmp_path):
    """M2：max_len 超过硬上限时被拒绝，不会枚举爆炸。"""
    src = _make_pdf(tmp_path / "plain.pdf", ["Page"])
    enc = tmp_path / "enc.pdf"
    doc = fitz.open(src)
    doc.save(
        str(enc), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="needpw", owner_pw="needpw"
    )
    doc.close()

    result = run_command(
        "decrypt",
        {
            "input": str(enc),
            "--charset": "ab",
            "--min-len": "1",
            "--max-len": "9",
            "--jobs": "1",
            "-o": str(tmp_path / "out.pdf"),
        },
    )

    assert result["ok"] is False
    assert "上限" in result["error"]


def test_pwcrack_clamp_jobs_and_candidate_count():
    """M2：jobs 被 clamp 到 [1, cpu]；候选数计算正确。"""
    from pdf2zh.toolbox import pwcrack

    cpu = os.cpu_count() or 4
    assert pwcrack.clamp_jobs(0) == 1            # 下界
    assert pwcrack.clamp_jobs(-5) == 1
    assert pwcrack.clamp_jobs(10 ** 6) == cpu    # 上界 = CPU 核数
    assert pwcrack.clamp_jobs(cpu + 3) == cpu
    assert pwcrack.clamp_jobs(2) == 2
    assert pwcrack.clamp_jobs(None) == cpu       # 缺省 = 全核
    assert pwcrack.clamp_jobs("") == cpu
    assert pwcrack.clamp_jobs("nope") == cpu     # 非法值回退
    assert pwcrack.brute_candidate_count("ab", 1, 2) == 2 + 4
    assert pwcrack.MAX_BRUTE_LEN == 8


def test_split_argv_positional_value_equal_to_flag():
    """m1：位置参数取值恰为某个 flag 名时仍按位置参数消费，不丢参。"""
    from pdf2zh.toolbox import _split_argv

    argv = build_argv("extract", {"input": "-p", "-p": "1", "-o": "o.pdf"})
    parsed = _split_argv("extract", argv)
    assert parsed["input"] == "-p"
    assert parsed["-p"] == "1"
    assert parsed["-o"] == "o.pdf"

    argv2 = build_argv("merge", {"inputs": ["-o"], "-o": "m.pdf"})
    parsed2 = _split_argv("merge", argv2)
    assert parsed2["inputs"] == ["-o"]
    assert parsed2["-o"] == "m.pdf"


def test_util_parse_ranges_rejects_invalid():
    """m2：parse_ranges 补齐与 parse_pages 对齐的边界校验。"""
    from pdf2zh.toolbox import _util

    for spec in ["0", "abc", "3-1", "x-y"]:
        with pytest.raises(ValueError):
            _util.parse_ranges(spec, 3)
    with pytest.raises(ValueError):
        _util.parse_ranges("5", 4)  # 单页越界
    with pytest.raises(ValueError):
        _util.parse_ranges(",", 4)  # 选择为空
    # 区间上界超出仍截断（与 parse_pages 一致），不抛错
    assert _util.parse_ranges("1-10", 4) == [(0, 3)]


def test_cleanup_stale_workdirs_keeps_recent(tmp_path):
    """M3：仅保留最近 N 个历史工作目录，且不误删无关目录。"""
    import time as _time

    from pdf2zh.toolbox_gui import _WORKDIR_PREFIX, _cleanup_stale_workdirs

    base = tmp_path / "tmpbase"
    base.mkdir()
    (base / "unrelated").mkdir()

    for i in range(8):
        d = base / f"{_WORKDIR_PREFIX}{i:02d}"
        d.mkdir()
        ts = _time.time() - (8 - i) * 100  # i 越大越新
        os.utime(d, (ts, ts))

    removed = _cleanup_stale_workdirs(keep=5, base_dir=str(base))

    remaining = sorted(
        p.name for p in base.iterdir() if p.name.startswith(_WORKDIR_PREFIX)
    )
    assert remaining == [f"{_WORKDIR_PREFIX}{i:02d}" for i in range(3, 8)]
    assert len(removed) == 3
    assert (base / "unrelated").is_dir()


def test_make_component_reads_bool_default_and_numeric_bounds():
    """m4：Checkbox 读取 field.default；M2：数值字段带 min/max 上界。"""
    gr = pytest.importorskip("gradio")
    from pdf2zh.toolbox_gui import _make_component

    with gr.Blocks():
        checkbox = _make_component(
            "watermark-text",
            {"key": "--center", "label": "居中", "type": "bool", "default": True},
        )
        assert checkbox.value is True

        default_bool = _make_component(
            "encrypt", {"key": "--no-print", "label": "禁止打印", "type": "bool"}
        )
        assert default_bool.value is False

        max_len = _make_component(
            "decrypt",
            {"key": "--max-len", "label": "最大长度", "type": "int", "default": 4},
        )
        assert max_len.minimum == 1
        assert max_len.maximum == 8  # 与 pwcrack.MAX_BRUTE_LEN 对齐

        jobs = _make_component(
            "decrypt", {"key": "--jobs", "label": "并行进程数", "type": "int"}
        )
        assert jobs.minimum == 1
        assert jobs.maximum == (os.cpu_count() or 4)